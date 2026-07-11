import html
import json
import re
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path("C:/Users/asus/Downloads/c\u00f4ng thuc_1.xlsx")
OUTPUT = Path(__file__).resolve().parents[1] / "index.html"


LABELS = {
    "hsd": "HSD",
    "cách chế biến": "Cách chế biến",
    "cách chế biến": "Cách chế biến",
    "chế biến": "Chế biến",
    "sơ chế": "Sơ chế",
    "công thức": "Công thức",
    "nguyên liệu": "Nguyên liệu",
    "định lượng": "Định lượng",
    "điều kiện bảo quản": "Điều kiện bảo quản",
    "thời gian hâm nóng": "Thời gian hâm nóng",
}


def norm(text):
    return re.sub(r"\s+", " ", str(text or "").strip()).lower().strip(":")


def is_probable_title(value):
    value = str(value or "").strip()
    if not value:
        return False
    key = norm(value)
    if key in LABELS:
        return False
    if re.search(r"\d+\s*(?:ml|g|kg|oz|p|phút|giây|s|cái|hộp|viên|cây)\b", value, flags=re.I):
        return False
    if ":" in value:
        return False
    first_word = key.split(" ", 1)[0]
    if first_word in {"cho", "lấy", "dùng", "đong", "đổ", "mở", "cắt", "vắt", "kẹp", "trưng", "hâm", "fill", "refill"}:
        return False
    letters = [c for c in value if c.isalpha()]
    if not letters:
        return False
    uppercase_ratio = sum(1 for c in letters if c.isupper()) / len(letters)
    if uppercase_ratio > 0.72:
        return True
    return len(value) <= 34 and "\n" not in value and len(value.split()) <= 5


def split_steps(value):
    parts = []
    for raw_line in str(value or "").splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if "→" in line or "->" in line or "=>" in line:
            chunks = re.split(r"\s*(?:→|->|=>)\s*", line)
            parts.extend(chunk.strip() for chunk in chunks if chunk.strip())
        else:
            parts.append(line)
    return parts


def tags_for_text(text):
    found = []
    text = str(text or "")
    patterns = [
        (r"\bHSD\b[^.\n;]*", "HSD"),
        (r"\bD\s*\+\s*\d+", "HSD"),
        (r"\bN\s*\+\s*\d+", "HSD"),
        (r"\d+\s*(?:p|phút|giây|s)\b", "Thời gian"),
        (r"\d+\s*độ", "Nhiệt độ"),
        (r"\d+\s*(?:ml|g|kg|oz)\b", "Định lượng"),
    ]
    for pattern, label in patterns:
        matches = re.findall(pattern, text, flags=re.I)
        for match in matches[:3]:
            tag = re.sub(r"\s+", " ", match).strip()
            if label == "HSD" and not tag.upper().startswith(("HSD", "D", "N")):
                tag = f"HSD: {tag}"
            if tag not in found:
                found.append(tag)
    return found[:8]


def cell_value(ws, row, col):
    value = ws.cell(row, col).value
    if value is None:
        return ""
    return str(value).strip()


def lane_pairs(max_col):
    pairs = []
    starts = [1, 4, 7, 10, 13, 15]
    for start in starts:
        if start <= max_col:
            pairs.append((start, min(start + 1, max_col)))
    return pairs


def add_entry(card, label, text, source):
    text = str(text or "").strip()
    if not text:
        return
    card["entries"].append(
        {
            "label": label or "Ghi chú",
            "text": text,
            "steps": split_steps(text),
            "source": source,
        }
    )
    for tag in tags_for_text(text):
        if tag not in card["tags"]:
            card["tags"].append(tag)


def parse_lane(ws, label_col, text_col):
    cards = []
    current = None
    last_label = None

    def ensure_card(title=None):
        nonlocal current
        if current is None:
            current = {
                "title": title or "Ghi chú chung",
                "entries": [],
                "tags": [],
                "source": f"{ws.title}",
            }
            cards.append(current)
        return current

    for row in range(1, ws.max_row + 1):
        left = cell_value(ws, row, label_col)
        right = cell_value(ws, row, text_col)
        if not left and not right:
            continue

        left_key = norm(left)
        right_is_title = right and not left and is_probable_title(right)
        left_is_title_only = left and not right and is_probable_title(left)

        if right_is_title:
            current = {
                "title": right,
                "entries": [],
                "tags": [],
                "source": f"{ws.title}!{ws.cell(row, text_col).coordinate}",
            }
            cards.append(current)
            last_label = None
            continue

        if left_is_title_only:
            current = {
                "title": left,
                "entries": [],
                "tags": [],
                "source": f"{ws.title}!{ws.cell(row, label_col).coordinate}",
            }
            cards.append(current)
            last_label = None
            continue

        if left and right and left_key not in LABELS and is_probable_title(left):
            current = {
                "title": left,
                "entries": [],
                "tags": [],
                "source": f"{ws.title}!{ws.cell(row, label_col).coordinate}",
            }
            cards.append(current)
            add_entry(current, "Công thức", right, ws.cell(row, text_col).coordinate)
            last_label = "Công thức"
            continue

        card = ensure_card()
        if left_key in LABELS:
            label = LABELS[left_key]
            last_label = label
            add_entry(card, label, right or left, ws.cell(row, text_col).coordinate)
        elif left and right:
            label = left.strip(":")
            last_label = label
            add_entry(card, label, right, ws.cell(row, text_col).coordinate)
        elif right:
            add_entry(card, last_label or "Bước", right, ws.cell(row, text_col).coordinate)
        elif left:
            add_entry(card, last_label or "Bước", left, ws.cell(row, label_col).coordinate)

    return [card for card in cards if card["entries"] or card["tags"]]


def parse_workbook(path):
    wb = load_workbook(path, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        sheet = {"name": ws.title, "cards": []}
        seen = set()
        for label_col, text_col in lane_pairs(ws.max_column):
            for card in parse_lane(ws, label_col, text_col):
                signature = (
                    card["title"],
                    tuple((entry["label"], entry["text"]) for entry in card["entries"]),
                )
                if signature in seen:
                    continue
                seen.add(signature)
                sheet["cards"].append(card)
        if sheet["cards"]:
            sheets.append(sheet)
    return sheets


def render_html(data):
    data_json = json.dumps(data, ensure_ascii=False)
    return f"""<!doctype html>
<html lang="vi">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Sổ tay công thức</title>
  <style>
    :root {{
      --bg: #f7f8fb;
      --surface: #ffffff;
      --surface-2: #eef2f7;
      --ink: #1c2430;
      --muted: #657386;
      --line: #d9e0ea;
      --primary: #0f766e;
      --primary-strong: #0b5f59;
      --accent: #b45309;
      --danger: #b42318;
      --shadow: 0 14px 35px rgba(28, 36, 48, .08);
    }}
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0;
      font-family: Arial, Helvetica, sans-serif;
      color: var(--ink);
      background: var(--bg);
    }}
    header {{
      background: var(--surface);
      border-bottom: 1px solid var(--line);
      position: sticky;
      top: 0;
      z-index: 3;
    }}
    .topbar {{
      max-width: 1180px;
      margin: 0 auto;
      padding: 18px 18px 14px;
    }}
    h1 {{
      margin: 0 0 12px;
      font-size: 28px;
      line-height: 1.2;
      font-weight: 700;
    }}
    .controls {{
      display: grid;
      grid-template-columns: minmax(220px, 1fr) 220px;
      gap: 10px;
      align-items: end;
    }}
    label {{
      display: block;
      font-size: 13px;
      color: var(--muted);
      margin-bottom: 5px;
      font-weight: 700;
    }}
    input, select {{
      width: 100%;
      min-height: 42px;
      border: 1px solid var(--line);
      border-radius: 8px;
      padding: 9px 11px;
      color: var(--ink);
      background: var(--surface);
      font-size: 15px;
    }}
    main {{
      max-width: 1180px;
      margin: 0 auto;
      padding: 18px;
    }}
    .stats {{
      display: grid;
      grid-template-columns: repeat(3, minmax(0, 1fr));
      gap: 10px;
      margin-bottom: 16px;
    }}
    .stat {{
      background: var(--surface);
      border: 1px solid var(--line);
      border-radius: 8px;
      padding: 12px;
    }}
    .stat span {{
      display: block;
      color: var(--muted);
      font-size: 13px;
      margin-bottom: 4px;
    }}
    .stat strong {{
      font-size: 22px;
    }}
    .tabs {{
      display: flex;
      flex-wrap: wrap;
      gap: 8px;
      margin-bottom: 18px;
    }}
    .tab {{
      border: 1px solid var(--line);
      background: var(--surface);
      border-radius: 999px;
      padding: 8px 12px;
      cursor: pointer;
      color: var(--ink);
      font-size: 14px;
    }}
    .tab.active {{
      background: var(--primary);
      border-color: var(--primary);
      color: #fff;
    }}
    .sheet-title {{
      display: flex;
      align-items: baseline;
      justify-content: space-between;
      gap: 12px;
      margin: 22px 0 12px;
    }}
    .sheet-title h2 {{
      margin: 0;
      font-size: 22px;
    }}
    .sheet-title span {{
      color: var(--muted);
      font-size: 13px;
    }}
    .grid {{
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(300px, 1fr));
      gap: 14px;
    }}
    .card {{
      background: var(--surface);
      border: 1px solid var(--line);
      border-radius: 8px;
      box-shadow: var(--shadow);
      overflow: hidden;
      display: flex;
      flex-direction: column;
      min-height: 100%;
    }}
    .card-head {{
      padding: 14px 14px 10px;
      border-bottom: 1px solid var(--line);
    }}
    .card h3 {{
      margin: 0 0 9px;
      font-size: 17px;
      line-height: 1.3;
    }}
    .tags {{
      display: flex;
      flex-wrap: wrap;
      gap: 6px;
    }}
    .tag {{
      color: var(--primary-strong);
      background: rgba(15, 118, 110, .1);
      border: 1px solid rgba(15, 118, 110, .18);
      border-radius: 999px;
      padding: 4px 8px;
      font-size: 12px;
      line-height: 1.2;
    }}
    .tag.hot {{ color: var(--danger); background: rgba(180,35,24,.09); border-color: rgba(180,35,24,.18); }}
    .tag.qty {{ color: var(--accent); background: rgba(180,83,9,.1); border-color: rgba(180,83,9,.18); }}
    .entries {{
      padding: 12px 14px 14px;
      display: grid;
      gap: 12px;
    }}
    .entry-label {{
      display: inline-flex;
      width: fit-content;
      margin-bottom: 7px;
      color: var(--muted);
      font-size: 13px;
      font-weight: 700;
    }}
    ol {{
      margin: 0;
      padding-left: 22px;
    }}
    li {{
      margin: 0 0 6px;
      line-height: 1.45;
    }}
    .source {{
      margin-top: 6px;
      color: var(--muted);
      font-size: 12px;
    }}
    .empty {{
      background: var(--surface);
      border: 1px dashed var(--line);
      border-radius: 8px;
      padding: 26px;
      color: var(--muted);
      text-align: center;
    }}
    mark {{
      background: rgba(180, 83, 9, .18);
      color: inherit;
      border-radius: 4px;
      padding: 0 2px;
    }}
    @media (max-width: 720px) {{
      .controls, .stats {{ grid-template-columns: 1fr; }}
      h1 {{ font-size: 24px; }}
      main, .topbar {{ padding-left: 12px; padding-right: 12px; }}
      .grid {{ grid-template-columns: 1fr; }}
    }}
  </style>
</head>
<body>
  <header>
    <div class="topbar">
      <h1>Sổ tay công thức pha chế & chế biến</h1>
      <div class="controls">
        <div>
          <label for="search">Tìm món, nguyên liệu, định lượng, HSD</label>
          <input id="search" type="search" placeholder="Ví dụ: matcha, 30ml, bánh bao, D+2">
        </div>
        <div>
          <label for="sheetSelect">Nhóm công thức</label>
          <select id="sheetSelect"></select>
        </div>
      </div>
    </div>
  </header>
  <main>
    <section class="stats" aria-label="Tổng quan">
      <div class="stat"><span>Nhóm</span><strong id="sheetCount">0</strong></div>
      <div class="stat"><span>Mục công thức</span><strong id="cardCount">0</strong></div>
      <div class="stat"><span>Đang hiển thị</span><strong id="visibleCount">0</strong></div>
    </section>
    <nav class="tabs" id="tabs" aria-label="Nhóm công thức"></nav>
    <section id="content"></section>
  </main>
  <script>
    const DATA = {data_json};
    const searchInput = document.getElementById('search');
    const sheetSelect = document.getElementById('sheetSelect');
    const tabs = document.getElementById('tabs');
    const content = document.getElementById('content');
    const sheetCount = document.getElementById('sheetCount');
    const cardCount = document.getElementById('cardCount');
    const visibleCount = document.getElementById('visibleCount');
    let activeSheet = 'all';

    const totalCards = DATA.reduce((sum, sheet) => sum + sheet.cards.length, 0);
    sheetCount.textContent = DATA.length;
    cardCount.textContent = totalCards;

    function escapeRegExp(value) {{
      return value.replace(/[.*+?^${{}}()|[\\]\\\\]/g, '\\\\$&');
    }}

    function highlight(text, query) {{
      const safe = String(text || '').replace(/[&<>"']/g, char => ({{
        '&': '&amp;', '<': '&lt;', '>': '&gt;', '"': '&quot;', "'": '&#039;'
      }}[char]));
      if (!query) return safe;
      return safe.replace(new RegExp(escapeRegExp(query), 'ig'), match => `<mark>${{match}}</mark>`);
    }}

    function plainCardText(card, sheetName) {{
      return [sheetName, card.title, ...card.tags, ...card.entries.flatMap(entry => [entry.label, entry.text])]
        .join(' ')
        .toLowerCase();
    }}

    function tagClass(tag) {{
      if (/độ|phút|giây|\\bs\\b|\\bp\\b/i.test(tag)) return 'hot';
      if (/ml|g|kg|oz/i.test(tag)) return 'qty';
      return '';
    }}

    function renderTabs() {{
      const options = [['all', 'Tất cả'], ...DATA.map(sheet => [sheet.name, sheet.name.replaceAll('_', ' ')])];
      sheetSelect.innerHTML = options.map(([value, label]) => `<option value="${{value}}">${{label}}</option>`).join('');
      tabs.innerHTML = options.map(([value, label]) => (
        `<button class="tab ${{value === activeSheet ? 'active' : ''}}" type="button" data-sheet="${{value}}">${{label}}</button>`
      )).join('');
      sheetSelect.value = activeSheet;
    }}

    function render() {{
      const query = searchInput.value.trim().toLowerCase();
      renderTabs();
      let visible = 0;
      const sheetHtml = DATA
        .filter(sheet => activeSheet === 'all' || sheet.name === activeSheet)
        .map(sheet => {{
          const cards = sheet.cards.filter(card => plainCardText(card, sheet.name).includes(query));
          visible += cards.length;
          if (!cards.length) return '';
          return `
            <section>
              <div class="sheet-title">
                <h2>${{highlight(sheet.name.replaceAll('_', ' '), query)}}</h2>
                <span>${{cards.length}} mục</span>
              </div>
              <div class="grid">
                ${{cards.map(card => `
                  <article class="card">
                    <div class="card-head">
                      <h3>${{highlight(card.title, query)}}</h3>
                      <div class="tags">
                        ${{card.tags.slice(0, 8).map(tag => `<span class="tag ${{tagClass(tag)}}">${{highlight(tag, query)}}</span>`).join('')}}
                      </div>
                    </div>
                    <div class="entries">
                      ${{card.entries.map(entry => `
                        <div class="entry">
                          <span class="entry-label">${{highlight(entry.label, query)}}</span>
                          <ol>
                            ${{entry.steps.map(step => `<li>${{highlight(step, query)}}</li>`).join('')}}
                          </ol>
                          <div class="source">${{highlight(sheet.name + ' · ' + entry.source, query)}}</div>
                        </div>
                      `).join('')}}
                    </div>
                  </article>
                `).join('')}}
              </div>
            </section>`;
        }}).join('');
      visibleCount.textContent = visible;
      content.innerHTML = sheetHtml || '<div class="empty">Không tìm thấy mục phù hợp.</div>';
    }}

    tabs.addEventListener('click', event => {{
      const button = event.target.closest('button[data-sheet]');
      if (!button) return;
      activeSheet = button.dataset.sheet;
      render();
    }});
    sheetSelect.addEventListener('change', event => {{
      activeSheet = event.target.value;
      render();
    }});
    searchInput.addEventListener('input', render);
    render();
  </script>
</body>
</html>
"""


def main():
    data = parse_workbook(SOURCE)
    OUTPUT.write_text(render_html(data), encoding="utf-8")
    print(f"Wrote {OUTPUT}")
    print(f"Sheets: {len(data)}")
    print(f"Cards: {sum(len(sheet['cards']) for sheet in data)}")


if __name__ == "__main__":
    main()
